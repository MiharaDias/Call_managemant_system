from datetime import timedelta
from pathlib import Path

from app.common import today
from app.services.phones import normalize_phone

FIXTURE = Path(__file__).parent / "fixtures" / "Teacher_Call_Tracker.xlsx"


def test_phone_normalisation():
    assert normalize_phone("071 344 4555") == "+94713444555"
    assert normalize_phone("+94 71 344 4555") == "+94713444555"
    assert normalize_phone("94713444555") == "+94713444555"
    assert normalize_phone("") == ""
    assert normalize_phone("12345") == ""


def test_routes_require_login(client):
    client.post("/api/auth/logout")
    for path in ("/api/teachers", "/api/calls", "/api/tasks", "/api/tasks/agenda", "/api/dashboard/summary",
                 "/api/options", "/api/export/excel", "/api/auth/me"):
        assert client.get(path).status_code == 401, path


def test_login_rejects_bad_password(client):
    assert client.post("/api/auth/login", json={"username": "alice", "password": "nope"}).status_code == 401
    assert client.post("/api/auth/login", json={"username": "ghost", "password": "nope"}).status_code == 401


def test_import_workbook_then_full_flow(authed):
    c = authed
    with open(FIXTURE, "rb") as f:
        r = c.post("/api/import/excel", files={"file": ("t.xlsx", f)})
    assert r.status_code == 200, r.text
    assert r.json()["teachers_created"] == 26
    # importing again does not duplicate
    with open(FIXTURE, "rb") as f:
        assert c.post("/api/import/excel", files={"file": ("t.xlsx", f)}).json()["teachers_created"] == 0

    teachers = c.get("/api/teachers").json()
    assert len(teachers) == 26
    assert len(c.get("/api/teachers", params={"has_phone": False}).json()) == 6  # the 6 "no number" teachers
    assert len(c.get("/api/teachers", params={"q": "0713444555"}).json()) == 1  # digits-only phone search
    hit = c.get("/api/teachers", params={"q": "071 344 4555"}).json()
    assert [t["name"] for t in hit] == ["Darshana Ukuwela"]
    assert {t["name"] for t in c.get("/api/teachers", params={"subject": "Econ"}).json()} == {
        "Lucky Chathuranga", "Kasun Liyanage"}
    assert len(c.get("/api/teachers", params={"rounds": 0}).json()) == 26

    tid = hit[0]["id"]
    today_ = today()
    base = {"teacher_id": tid, "called_date": today_.isoformat(), "called_time": "10:30:00"}

    # Round 1: no answer, then a second attempt in the same round that gets answered
    r = c.post("/api/calls", json={**base, "round": 1, "outcome": "No answer"})
    assert r.status_code == 201 and r.json()["attempt"] == 1
    r = c.post("/api/calls", json={**base, "round": 1, "outcome": "Answered - Spoke", "spoke_to": "Manager",
                                   "manager_contacted": True, "demo_sent": True, "price_proposed": 15000,
                                   "interest_level": "Considering"})
    assert r.json()["attempt"] == 2
    # Round 2 with a follow-up task created in the same request
    tomorrow = today_ + timedelta(days=1)
    r = c.post("/api/calls", json={
        **base, "round": 2, "outcome": "Answered - Spoke", "spoke_to": "Teacher", "price_proposed": 12000,
        "interest_level": "Negotiating", "next_action": "Call again", "followup_date": tomorrow.isoformat(),
        "create_tasks": [{"type": "Call back", "title": "Call again", "due_date": tomorrow.isoformat(),
                          "due_time": "09:00:00"}],
    })
    assert r.status_code == 201

    t = c.get(f"/api/teachers/{tid}").json()
    assert t["rounds"]["1"] == {"outcome": "Answered - Spoke", "attempts": 2}
    assert t["rounds_attempted"] == 2 and t["next_round"] == 3 and t["call_count"] == 3
    assert t["reached"] and t["manager_contacted"] and t["demo_sent"] and not t["detailed_sent"]
    assert t["latest_price"] == 12000 and t["latest_interest"] == "Negotiating"
    assert t["status"] == "Negotiating"
    assert t["next_followup"] == tomorrow.isoformat()
    assert len(t["calls"]) == 3

    # Agenda: the task is on tomorrow, not today
    ag = c.get("/api/tasks/agenda", params={"date": tomorrow.isoformat()}).json()
    assert [x["title"] for x in ag["tasks"]] == ["Call again"]
    assert c.get("/api/tasks/agenda").json()["tasks"] == []

    # An overdue task shows up on today's agenda, and marking it done removes it from overdue
    r = c.post("/api/tasks", json={"teacher_id": tid, "type": "Send demo video", "title": "Send video",
                                   "due_date": (today_ - timedelta(days=2)).isoformat()})
    task_id = r.json()["id"]
    assert r.json()["overdue"] is True
    ag = c.get("/api/tasks/agenda").json()
    assert [x["id"] for x in ag["overdue"]] == [task_id]
    done = c.patch(f"/api/tasks/{task_id}", json={"status": "done"}).json()
    assert done["status"] == "done" and done["completed_by"] == "Alice" and not done["overdue"]
    assert c.get("/api/tasks/agenda").json()["overdue"] == []
    assert c.patch(f"/api/tasks/{task_id}", json={"status": "open"}).json()["completed_by"] == ""

    # Calendar range
    cal = c.get("/api/tasks/calendar", params={"start": today_.isoformat(), "end": tomorrow.isoformat()}).json()
    assert any(x["title"] == "Call again" for x in cal)

    # Dashboard mirrors the Excel formulas
    d = c.get("/api/dashboard/summary").json()
    rows = {r["key"]: r for r in d["table"]}
    assert rows["calls_logged"]["rounds"]["1"] == 2 and rows["calls_logged"]["total"] == 3
    assert rows["answered"]["total"] == 2 and rows["no_contact"]["total"] == 1
    assert rows["spoke_teacher"]["total"] == 1 and rows["manager_contacted"]["total"] == 1
    assert rows["prices_count"]["total"] == 2 and rows["avg_price"]["total"] == 13500
    assert d["totals"]["teachers"] == 26 and d["totals"]["never_called"] == 25
    assert d["by_status"]["Negotiating"] == 1
    assert d["tasks_overdue"] == 1  # the re-opened task
    assert all(q["teacher"]["id"] != tid for q in d["next_to_call"])  # has an open task already

    # Export round trip produces a readable workbook
    from io import BytesIO
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(c.get("/api/export/excel").content))
    assert wb.sheetnames == ["Contacts", "Call Log", "Tasks"]
    assert wb["Contacts"].max_row == 27 and wb["Call Log"].max_row == 4

    # Deleting a teacher removes calls and tasks with it
    assert c.delete(f"/api/teachers/{tid}").status_code == 204
    assert c.get(f"/api/teachers/{tid}").status_code == 404
    assert c.get("/api/calls").json() == []


def test_rounds_exhausted_and_options(authed):
    c = authed
    t = c.post("/api/teachers", json={"name": "Test Teacher", "phones": [{"number": "077 111 2222", "label": "Hotline"}]}).json()
    for rnd in range(1, 6):
        c.post("/api/calls", json={"teacher_id": t["id"], "round": rnd, "called_date": today().isoformat(),
                                   "outcome": "No answer"})
    t = c.get(f"/api/teachers/{t['id']}").json()
    assert t["rounds_exhausted"] and t["next_round"] == 6  # extra rounds are allowed

    opts = c.get("/api/options").json()
    assert "Answered - Spoke" in opts["options"]["outcome"] and opts["max_rounds"] == 5
    opts["options"]["channel"].append("Telegram")
    assert "Telegram" in c.put("/api/options", json=opts).json()["options"]["channel"]


def test_hotline_flow_extra_rounds_tags_and_contacted(authed):
    c = authed
    t = c.post("/api/teachers", json={"name": "Hotline Teacher", "subject": "Astronomy",
                                      "phones": [{"number": "011 222 3333", "label": "Hotline"}]}).json()
    tid = t["id"]
    # a brand-new subject typed on the form becomes a selectable option
    assert "Astronomy" in c.get("/api/options").json()["options"]["subject"]
    # status list has Contacted straight after Not started
    statuses = c.get("/api/options").json()["options"]["status"]
    assert statuses[:2] == ["Not started", "Contacted"]

    # tags: create, attach to teacher + call, filter both lists
    tag = c.post("/api/tags", json={"name": "Decision maker", "color": "#16a34a", "description": "Can approve the deal"}).json()
    assert c.post("/api/tags", json={"name": "decision MAKER"}).status_code == 409
    c.patch(f"/api/teachers/{tid}", json={"tag_ids": [tag["id"]]})
    assert [x["name"] for x in c.get("/api/teachers", params={"tag_ids": tag["id"]}).json()] == ["Hotline Teacher"]
    assert c.get("/api/teachers", params={"tag_ids": tag["id"] + 99}).json() == []

    # 1) call the hotline: they give us the manager's number. Spoke to hotline is not yet "Contacted".
    r = c.post("/api/calls", json={"teacher_id": tid, "round": 1, "called_date": today().isoformat(),
                                   "outcome": "Answered - Spoke", "spoke_to": "Hotline",
                                   "new_phone": {"number": "077 999 8888", "label": "Manager"}}).json()
    t = c.get(f"/api/teachers/{tid}").json()
    assert [(p["number"], p["label"]) for p in t["phones"]] == [("011 222 3333", "Hotline"), ("077 999 8888", "Manager")]
    assert t["status"] == "In progress"
    # the same number is not added twice
    c.post("/api/calls", json={"teacher_id": tid, "round": 1, "called_date": today().isoformat(), "outcome": "Busy",
                               "new_phone": {"number": "0779998888"}})
    assert len(c.get(f"/api/teachers/{tid}").json()["phones"]) == 2

    # 2) call the manager -> Contacted; tag this call
    c.post("/api/calls", json={"teacher_id": tid, "round": 2, "called_date": today().isoformat(),
                               "outcome": "Answered - Spoke", "spoke_to": "Manager", "tag_ids": [tag["id"]]})
    assert c.get(f"/api/teachers/{tid}").json()["status"] == "Contacted"
    calls = c.get("/api/calls", params={"tag_ids": tag["id"]}).json()
    assert len(calls) == 3  # one tagged call + the teacher carries the tag so all its calls match
    other = c.post("/api/teachers", json={"name": "Other"}).json()
    c.post("/api/calls", json={"teacher_id": other["id"], "round": 1, "called_date": today().isoformat(), "outcome": "Busy"})
    assert all(x["teacher_name"] == "Hotline Teacher" for x in calls)

    # 3) rounds beyond the default of five
    for rnd in (3, 4, 5, 6, 7):
        c.post("/api/calls", json={"teacher_id": tid, "round": rnd, "called_date": today().isoformat(), "outcome": "No answer"})
    t = c.get(f"/api/teachers/{tid}").json()
    assert t["rounds_attempted"] == 7 and t["next_round"] == 8 and not t["rounds_exhausted"]  # reached us earlier
    d = c.get("/api/dashboard/summary").json()
    assert d["rounds"] == [1, 2, 3, 4, 5, 6, 7]

    # custom status can be added from a form and used
    assert "Waiting for approval" in c.post("/api/options/status", json={"value": "Waiting for approval"}).json()["options"]["status"]
    assert c.patch(f"/api/teachers/{tid}", json={"status": "Waiting for approval"}).json()["status"] == "Waiting for approval"

    # tag deletion detaches it everywhere
    assert c.delete(f"/api/tags/{tag['id']}").status_code == 204
    assert c.get(f"/api/teachers/{tid}").json()["tags"] == []


def test_calendar_shows_only_selected_task_types(authed):
    c = authed
    day = today().isoformat()
    for typ, title in (("Meeting", "Meet at office"), ("Call back", "Ring back"), ("Send demo video", "Send clip")):
        c.post("/api/tasks", json={"type": typ, "title": title, "due_date": day})
    titles = {x["title"] for x in c.get("/api/tasks/calendar", params={"start": day, "end": day}).json()}
    assert {"Meet at office", "Ring back"} <= titles and "Send clip" not in titles
    # everything still shows in the agenda
    assert "Send clip" in {x["title"] for x in c.get("/api/tasks/agenda").json()["tasks"]}
    # Settings can add a type to the calendar
    opts = c.get("/api/options").json()
    c.put("/api/options", json={"options": opts["options"], "calendar_types": ["Meeting", "Call back", "Send demo video"]})
    titles = {x["title"] for x in c.get("/api/tasks/calendar", params={"start": day, "end": day}).json()}
    assert "Send clip" in titles
    c.put("/api/options", json={"options": opts["options"], "calendar_types": ["Call back", "Meeting"]})


def test_class_size_and_papers_per_month(authed):
    c = authed
    t = c.post("/api/teachers", json={"name": "Sizes", "students_per_class": 45, "papers_per_month": 300}).json()
    assert (t["students_per_class"], t["papers_per_month"]) == (45, 300)
    t = c.patch(f"/api/teachers/{t['id']}", json={"papers_per_month": 450}).json()
    assert (t["students_per_class"], t["papers_per_month"]) == (45, 450)  # other field untouched
    t = c.patch(f"/api/teachers/{t['id']}", json={"students_per_class": None}).json()
    assert t["students_per_class"] is None and t["papers_per_month"] == 450  # can be cleared
    assert c.post("/api/teachers", json={"name": "Bad", "papers_per_month": -5}).status_code == 422
    from io import BytesIO
    from openpyxl import load_workbook
    ws = load_workbook(BytesIO(c.get("/api/export/excel").content))["Contacts"]
    assert [x.value for x in ws[1]][-2:] == ["Students per Class", "Papers per Month"]


def test_multiple_and_custom_next_actions(authed):
    c = authed
    t = c.post("/api/teachers", json={"name": "Multi Action"}).json()
    day, later = today().isoformat(), (today() + timedelta(days=3)).isoformat()
    assert "Send brochure" in c.post("/api/options/next_action", json={"value": "Send brochure"}).json()["options"]["next_action"]
    call = c.post("/api/calls", json={
        "teacher_id": t["id"], "round": 1, "called_date": day, "outcome": "Answered - Spoke", "spoke_to": "Teacher",
        "next_actions": ["Send demo video", "Meeting", "Send brochure", "Meeting"], "followup_date": day,
        "create_tasks": [
            {"type": "Send demo video", "title": "Send demo video", "due_date": day},
            {"type": "Meeting", "title": "Meeting", "due_date": later, "due_time": "10:30:00"},
            {"type": "Other", "title": "Send brochure", "due_date": later},
        ],
    }).json()
    assert call["next_actions"] == ["Send demo video", "Meeting", "Send brochure"]  # duplicates dropped, order kept
    assert call["next_action"] == "Send demo video, Meeting, Send brochure"
    tasks = c.get("/api/tasks", params={"teacher_id": t["id"]}).json()
    assert sorted(x["title"] for x in tasks) == ["Meeting", "Send brochure", "Send demo video"]
    assert all(x["assigned_to"] == "Alice" for x in tasks)
    # editing the actions
    edited = c.patch(f"/api/calls/{call['id']}", json={"next_actions": ["Call again"]}).json()
    assert edited["next_actions"] == ["Call again"]
    # the calendar shows the meeting but not the brochure
    cal = {x["title"] for x in c.get("/api/tasks/calendar", params={"start": later, "end": later}).json()}
    assert "Meeting" in cal and "Send brochure" not in cal


def test_log_physical_and_online_meetings(authed):
    c = authed
    day = today().isoformat()
    t = c.post("/api/teachers", json={"name": "Meeting Person", "phones": [{"number": "071 000 1111"}]}).json()
    tid = t["id"]
    assert "Held" in c.get("/api/options").json()["options"]["meeting_outcome"]

    # a call needs a round, a meeting does not
    assert c.post("/api/calls", json={"teacher_id": tid, "called_date": day, "outcome": "Busy"}).status_code == 422
    m = c.post("/api/calls", json={
        "teacher_id": tid, "kind": "physical_meeting", "called_date": day, "called_time": "10:00:00",
        "location": "Colombo office", "outcome": "Held", "spoke_to": "Manager", "price_proposed": 18000,
        "interest_level": "Interested", "next_actions": ["Send detailed message"],
    })
    assert m.status_code == 201, m.text
    m = m.json()
    assert (m["kind"], m["round"], m["attempt"], m["location"]) == ("physical_meeting", 0, 1, "Colombo office")
    o = c.post("/api/calls", json={"teacher_id": tid, "kind": "online_meeting", "called_date": day,
                                   "location": "https://meet.example/abc", "outcome": "Held", "spoke_to": "Teacher"}).json()
    assert o["kind"] == "online_meeting" and o["attempt"] == 1

    # meetings do not count as call rounds, but they do count as contact and carry price/interest
    v = c.get(f"/api/teachers/{tid}").json()
    assert v["rounds_attempted"] == 0 and v["call_count"] == 0 and v["meeting_count"] == 2
    assert v["next_round"] == 1 and v["reached"] and v["latest_price"] == 18000
    assert v["status"] == "Interested"
    assert len(v["calls"]) == 2

    # the first call after the meetings is still Round 1
    first = c.post("/api/calls", json={"teacher_id": tid, "round": 1, "called_date": day, "outcome": "No answer"}).json()
    assert first["round"] == 1 and first["attempt"] == 1 and first["kind"] == "call"

    assert {x["kind"] for x in c.get("/api/calls", params={"teacher_id": tid}).json()} == {"call", "physical_meeting", "online_meeting"}
    meetings = c.get("/api/calls", params={"teacher_id": tid, "kind": "physical_meeting"}).json()
    assert [x["location"] for x in meetings] == ["Colombo office"]

    # dashboard per-round table only counts calls; meetings are reported separately
    d = c.get("/api/dashboard/summary").json()
    assert d["totals"]["meetings"] >= 2
    logged = {r["key"]: r for r in d["table"]}["calls_logged"]
    assert logged["rounds"]["1"] >= 1 and 0 not in [int(k) for k in logged["rounds"]]

    from io import BytesIO
    from openpyxl import load_workbook
    ws = load_workbook(BytesIO(c.get("/api/export/excel").content))["Call Log"]
    types = [row[-2].value for row in ws.iter_rows(min_row=2)]
    assert "Physical meeting" in types and "Online meeting" in types and "Call" in types


def test_older_clients_single_create_task_still_creates_the_task(authed):
    """A browser tab left open on an older version sends `create_task` + `next_action`; the task must not be lost."""
    c = authed
    t = c.post("/api/teachers", json={"name": "Stale Tab"}).json()
    tomorrow = (today() + timedelta(days=1)).isoformat()
    call = c.post("/api/calls", json={
        "teacher_id": t["id"], "round": 1, "called_date": today().isoformat(), "outcome": "Answered - Spoke",
        "next_action": "Call again", "followup_date": tomorrow,
        "create_task": {"type": "Call back", "title": "Call again", "due_date": tomorrow},
    }).json()
    assert call["next_actions"] == ["Call again"]
    tasks = c.get("/api/tasks", params={"teacher_id": t["id"]}).json()
    assert [(x["title"], x["type"], x["due_date"], x["source_call_id"]) for x in tasks] == [("Call again", "Call back", tomorrow, call["id"])]
